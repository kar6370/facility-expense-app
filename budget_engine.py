"""Integer-won budget reconciliation; no UI or cloud dependencies.

Official payments are ledger snapshots. Transaction allocations explain those
payments, never add to them. All mutations are committed by BudgetStore.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import io
import json
import re
import sqlite3
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime
from calendar import monthrange
from decimal import Decimal, InvalidOperation
from pathlib import Path
from account_rules import source_account

SCHEMA = 44
STAGES = ["본예산", "제1회 추경", "제2회 추경"]
SCOPES = ["운영경비", "일상경비"]
MODES = {"auto": "자동 인식 · ERP 혼합 양식 포함", "full": "경정예산 전체명세", "partial": "변경 항목의 경정액", "delta": "변경 항목의 증감액"}
CATEGORY_CODES = {
    "전기요금": "215", "상하수도": "206-01", "일반재료비": "206-01",
    "통신요금": "201-21", "복합기임대": "201-13", "공청기비데": "201-11",
    "세탁용역": "201-11", "환경용역": "201-11", "무인경비": "201-11",
    "야간경비": "201-11", "승강기점검": "201-11", "신용카드수수료": "201-11",
    "부서업무비": "217-02", "상품매입비": "233", "자체소수선": "214-05",
    "수탁자산취득비": "405-12", "전산개발비": "207-02",
}


class ValidationError(ValueError):
    pass


class ConflictError(RuntimeError):
    pass


def norm(value):
    return re.sub(r"\s+", "", str(value or "")).lower()


def won(value, scale=1):
    if value is None or str(value).strip() in ("", "-", "#cspan", "#rspan"):
        return 0
    s = str(value).strip().replace(",", "").replace("원", "").replace(" ", "")
    negative = s.startswith(("△", "▲", "("))
    s = s.replace("△", "").replace("▲", "").strip("()").replace("−", "-")
    try:
        x = Decimal(s) * Decimal(scale)
    except InvalidOperation as exc:
        raise ValidationError(f"금액으로 읽을 수 없습니다: {value}") from exc
    if not x.is_finite() or x != x.to_integral_value():
        raise ValidationError(f"원 단위 정수가 아닙니다: {value}")
    return -abs(int(x)) if negative else int(x)


def parse_date(value):
    if isinstance(value, (date, datetime)):
        return value.date() if isinstance(value, datetime) else value
    s = str(value or "").strip()
    match = re.match(r"^(\d{4})[./년-]?\s*(\d{1,2})[./월-]?\s*(\d{1,2})(?:일|\b)", s)
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            return None
    return None


def transaction_signature(row):
    """Editable financial content, independent of import filename or order."""
    fields = [str(row.get(k, "")) for k in ("집행일자", "적요", "세목", "예산과목", "사업구분", "거래처", "거래상태", "일자근거")]
    # I-column changes can invalidate a previously confirmed leaf assignment.
    if row.get("원본I열"):
        fields.append(str(row["원본I열"]))
    fields.append(won(row.get("집행금액", 0)))
    return digest(json.dumps(fields, ensure_ascii=False).encode())


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def leaf_key(scope, code, path):
    # Preserve punctuation and the full ancestry; labels are not global IDs.
    normalized = [re.sub(r"\s+", " ", p).strip() for p in path]
    return json.dumps([scope, code, normalized], ensure_ascii=False, separators=(",", ":"))


def read_sheets(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        result = {}
        for ws in wb.worksheets:
            ws.reset_dimensions()  # ERP exports incorrectly declare dimension A1.
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            if any(any(v is not None for v in row) for row in rows):
                result[ws.title] = rows
        wb.close()
    except Exception as exc:
        raise ValidationError(f"엑셀을 열지 못했습니다: {exc}") from exc
    if not result:
        raise ValidationError("값이 있는 시트가 없습니다.")
    return result


def _scope(text, current):
    t = norm(text)
    if "운영" in t and "관리" in t and ("일상경비" in t or "수탁_일상" in t):
        return "일상경비"
    if "운영" in t and "관리" in t and not re.match(r"^\d", t):
        return "운영경비"
    return current


def inspect_workbook(raw, kind):
    sheets = read_sheets(raw)
    candidates = []
    for name, rows in sheets.items():
        for i, row in enumerate(rows[:30]):
            values = [norm(v) for v in row]
            if kind == "budget" and any("세목" in v for v in values) and any("예산액" in v or "증감액" in v for v in values):
                columns = {j: str(v) for j, v in enumerate(row) if v and j and "span" not in str(v)}
                candidates.append({"sheet": name, "header": i, "columns": columns})
                break
            if kind == "ledger" and any("지출액" in v for v in values) and any("예산배정" in v for v in values):
                candidates.append({"sheet": name, "header": i, "columns": {j: str(v) for j, v in enumerate(row) if v}})
                break
    if not candidates:
        raise ValidationError("필요한 제목행을 찾지 못했습니다. 예산명세서/통제원장 구분과 시트를 확인해 주세요.")
    return candidates


def parse_budget(raw, sheet=None, amount_column=None, scale=1000, mode="auto"):
    candidates = inspect_workbook(raw, "budget")
    selected = next((x for x in candidates if x["sheet"] == sheet), candidates[0])
    rows = read_sheets(raw)[selected["sheet"]]
    cols = selected["columns"]
    if amount_column is None:
        preferred = [i for i, v in cols.items() if "예산액" in norm(v) and not any(x in norm(v) for x in ("전년도", "기정", "증감"))]
        if mode == "delta":
            preferred = [i for i, v in cols.items() if "증감" in norm(v) and "비교" not in norm(v)]
        if not preferred:
            raise ValidationError("반영할 금액 열을 직접 선택해 주세요.")
        amount_column = preferred[0]
    if "전년도" in norm(cols.get(amount_column, "")):
        raise ValidationError("전년도 예산액은 이번 예산이나 추경 증감액으로 반영할 수 없습니다.")
    if mode == "delta" and "비교증감" in norm(cols.get(amount_column, "")) and any("전년도" in norm(v) for v in cols.values()):
        raise ValidationError("전년도 대비 비교증감을 추경 증감으로 사용할 수 없습니다.")
    requested_mode = mode
    mode = "full" if mode == "auto" else mode
    scope, mok, mok_name, code, name = SCOPES[0], "", "", "", ""
    path, levels, semoks, leaves, nodes, errors = [], [], {}, {}, [], []
    source_total = None
    previous_column = next((i for i, v in cols.items() if "전년도" in norm(v) or "기정" in norm(v)), None)
    delta_column = next((i for i, v in cols.items() if "증감" in norm(v)), None)
    def comparison(row):
        if previous_column is None or delta_column is None:
            return {}
        if max(previous_column, delta_column) >= len(row) or any(row[c] in (None, "") for c in (previous_column, delta_column)):
            return {}
        return {"previous_amount": won(row[previous_column], scale), "delta_amount": won(row[delta_column], scale)}
    total_comparison = {}
    for idx, row in enumerate(rows[selected["header"] + 1:], selected["header"] + 2):
        text = str(row[0] or "").strip() if row else ""
        if not text:
            continue
        val = row[amount_column] if amount_column < len(row) else None
        try:
            amount = won(val, scale)
        except ValidationError as exc:
            errors.append(f"{idx}행: {exc}"); continue
        changed_scope = _scope(text, scope)
        if "운영" in text and "관리" in text and not re.match(r"^\d", text) and not text.startswith(("◎", "○", "―")):
            scope, mok, code, path, levels = changed_scope, "", "", [], []
            continue
        if source_total is None and idx <= selected["header"] + 5 and val not in (None, ""):
            source_total = amount
            total_comparison = comparison(row)
        sm = re.match(r"^(\d{3}-\d{2})\.?\s+(.+)$", text)
        major = re.match(r"^(\d{3})\.?\s+(.+)$", text)
        if sm or (major and major.group(1) == mok and not code):
            m = sm or major
            code, name, path, levels = m.group(1), m.group(2).strip(), [], []
            key = scope + "|" + code
            if key in semoks:
                errors.append(f"{idx}행: 동일 사업구분·세목이 반복됩니다({key}).")
            semoks[key] = {"scope": scope, "code": code, "name": name, "mok": mok or code[:3], "mok_name": mok_name,
                           "amount": amount, "row": idx, "sheet": selected["sheet"], **comparison(row)}
            continue
        if major:
            mok, mok_name, code, path, levels = major.group(1), major.group(2).strip(), "", [], []
            continue
        marker = re.match(r"^([◎○◦―–—−-])\s*(.+)$", text)
        if marker and code:
            if val in (None, ""):
                errors.append(f"{idx}행: 부기 금액이 비어 있습니다. 수식 결과가 저장됐는지 확인해 주세요.")
            symbol, label = marker.groups()
            level = 1 if symbol == "◎" else 2 if symbol in ("○", "◦") else 3
            while levels and levels[-1] >= level:
                levels.pop(); path.pop()
            path.append(label.strip()); levels.append(level)
            key = leaf_key(scope, code, path)
            nodes.append({"id": key, "scope": scope, "code": code, "name": name, "mok": mok, "mok_name": mok_name,
                          "path": path.copy(), "label": " › ".join(path), "amount": amount, "row": idx,
                          "sheet": selected["sheet"], "basis": str(row[1] or "") if len(row) > 1 else ""})
    terminal_nodes = []
    for node in nodes:
        p = node["path"]
        children = [q for q in nodes if q["scope"] == node["scope"] and q["code"] == node["code"] and len(q["path"]) > len(p) and q["path"][:len(p)] == p]
        if not children:
            terminal_nodes.append(node)
    # ERP amendment exports retain current/prior totals on summary rows but
    # put signed changes in the budget column of detail rows. Header text alone
    # (often still '전년도예산액') cannot determine the detail's meaning.
    mixed = bool(total_comparison and terminal_nodes and semoks and
                 source_total - total_comparison["previous_amount"] == total_comparison["delta_amount"] and
                 all("delta_amount" in s and s["amount"] - s["previous_amount"] == s["delta_amount"] for s in semoks.values()) and
                 any(sum(n["amount"] for n in terminal_nodes if n["scope"]+"|"+n["code"] == k) == s["delta_amount"] != s["amount"]
                     for k, s in semoks.items()))
    if mixed:
        if requested_mode not in ("auto", "delta"):
            errors.append("세목은 경정 총액, 부기는 증감액인 ERP 추경 양식입니다. ‘자동 인식’을 선택하세요.")
        mode = "delta"
    for node in terminal_nodes:
        if node["id"] in leaves and not mixed:
            errors.append(f"{node['row']}행: 동일 경로의 부기가 반복됩니다({node['label']}).")
        if mixed:
            component = {"row": node["row"], "amount": node["amount"], "basis": node["basis"], "path": node["path"]}
            if node["id"] in leaves:
                leaves[node["id"]]["amount"] += node["amount"]
                leaves[node["id"]]["components"].append(component)
            else:
                leaves[node["id"]] = {**node, "components": [component]}
        else:
            leaves[node["id"]] = node
    if mixed:
        for key, semok in semoks.items():
            change = sum(x["amount"] for x in leaves.values() if x["scope"]+"|"+x["code"] == key)
            if change != semok["delta_amount"]:
                errors.append(f"{key}: 부기 증감 합계와 세목 증감액이 {change-semok['delta_amount']:+,}원 다릅니다.")
        if sum(x["amount"] for x in leaves.values()) != total_comparison["delta_amount"]:
            errors.append("부기 증감 합계가 명세서 전체 증감액과 다릅니다.")
    if not semoks or not leaves:
        errors.append("세목과 마지막 부기를 모두 읽지 못했습니다.")
    if mode == "full":
        for key, semok in semoks.items():
            amount = sum(x["amount"] for x in leaves.values() if x["scope"] + "|" + x["code"] == key)
            if amount != semok["amount"]:
                errors.append(f"{key}: 마지막 부기 합계와 세목 예산액이 {amount-semok['amount']:+,}원 다릅니다.")
        total = sum(s["amount"] for s in semoks.values())
        if source_total is None or total != source_total:
            errors.append(f"명세서 총액과 세목 합계가 다릅니다(원본 {source_total}, 인식 {total:,}).")
    if mode != "delta" and any(x["amount"] < 0 for x in leaves.values()):
        errors.append("경정액에 음수가 있습니다. 증감액 자료인지 확인해 주세요.")
    return {"kind": "budget", "mode": mode, "sheet": selected["sheet"], "column": amount_column,
            "column_label": cols.get(amount_column), "scale": scale, "source_total": source_total,
            "semoks": semoks, "leaves": leaves, "total": sum(x["amount"] for x in leaves.values()),
            "fingerprint": digest(raw), "errors": errors,
            "layout": "erp_mixed_delta" if mixed else "standard", **total_comparison}


def parse_ledger(raw, sheet=None, scale=1):
    choices = inspect_workbook(raw, "ledger")
    choice = next((x for x in choices if x["sheet"] == sheet), choices[0])
    rows = read_sheets(raw)[choice["sheet"]]
    h = choice["header"]
    top, sub = [norm(x) for x in rows[h]], [norm(x) for x in rows[h+1]]
    def col(values, predicate):
        return next((i for i, x in enumerate(values) if predicate(x)), None)
    indices = {
        "original": col(sub, lambda x: x == "당초"), "adjustment": col(sub, lambda x: x == "경정"),
        "carryover": col(sub, lambda x: x == "이월"), "transfer": col(sub, lambda x: x == "변경"),
        "expense_transfer": col(sub, lambda x: x == "변경(경비)"), "current_budget": col(sub, lambda x: x == "계"),
        "allocated": col(top, lambda x: "예산배정" in x), "expense": col(top, lambda x: "지출액" in x),
        "commitment": col(sub, lambda x: "원인행위(c)" in x),
    }
    if any(indices[x] is None for x in ("current_budget", "allocated", "expense")):
        raise ValidationError("예산현액(계), 예산배정액, 지출액(E) 열이 필요합니다.")
    scope, candidates, errors, total = SCOPES[0], {}, [], None
    for idx, row in enumerate(rows[h+2:], h+3):
        raw_label = str(row[0] or "") if row else ""
        label = raw_label.strip()
        if not label:
            continue
        scope = _scope(label, scope)
        try:
            vals = {k: won(row[i], scale) if i is not None and i < len(row) else 0 for k, i in indices.items()}
        except ValidationError as exc:
            errors.append(f"{idx}행: {exc}"); continue
        if label == "합계":
            total = vals; continue
        m = re.match(r"^(\d{3}(?:-\d{2})?)\.?\s+(.+)$", label)
        if m:
            code, name = m.groups()
            depth = len(raw_label.expandtabs()) - len(raw_label.expandtabs().lstrip())
            key = scope + "|" + code
            if key in candidates and candidates[key]["depth"] == depth:
                errors.append(f"{idx}행: 동일 수준의 세목이 반복됩니다({key}).")
            if key not in candidates or candidates[key]["depth"] <= depth:
                candidates[key] = {**vals, "scope": scope, "code": code, "name": name, "row": idx,
                                   "sheet": choice["sheet"], "depth": depth}
    lines = {k: v for k, v in candidates.items() if "-" in v["code"] or not any(k2.startswith(k+"-") for k2 in candidates)}
    if not total:
        errors.append("原本 합계 행을 찾지 못했습니다.".replace("原本", "원본"))
    else:
        for field in ("current_budget", "allocated", "expense"):
            delta = sum(x[field] for x in lines.values()) - total[field]
            if delta:
                errors.append(f"{field}: 원본 합계와 세목 합계가 {delta:+,}원 다릅니다.")
    return {"kind": "ledger", "lines": lines, "total": total or {}, "errors": errors,
            "sheet": choice["sheet"], "columns": indices, "scale": scale, "fingerprint": digest(raw)}


def empty_state(year):
    return {"schema": SCHEMA, "year": int(year), "stages": {}, "ledger": None, "assignments": {}, "overrides": {}}


def validate_state(state):
    if not isinstance(state, dict) or state.get("schema") != SCHEMA or type(state.get("year")) is not int or not 2000 <= state["year"] <= 2100:
        raise ValidationError("올바른 V44 회계연도 자료가 아닙니다.")
    if not all(isinstance(state.get(k), dict) for k in ("stages", "assignments", "overrides")):
        raise ValidationError("예산 차수·분류 자료 구조가 올바르지 않습니다.")
    for label, doc in state["stages"].items():
        if label not in STAGES or not isinstance(doc, dict) or doc.get("mode") not in ("full", "partial", "delta"):
            raise ValidationError("예산 차수 또는 반영 방식이 올바르지 않습니다.")
        if not isinstance(doc.get("leaves"), dict) or not isinstance(doc.get("semoks"), dict):
            raise ValidationError("세목·부기 자료가 없습니다.")
        for key, x in doc["leaves"].items():
            if not isinstance(x, dict) or x.get("scope") not in SCOPES or not isinstance(x.get("path"), list) or type(x.get("amount")) is not int:
                raise ValidationError("부기의 사업구분·경로·원 단위 금액이 올바르지 않습니다.")
            if not all(k in x for k in ("code", "id", "label", "row", "sheet")) or x["id"] != key:
                raise ValidationError("부기의 식별정보가 올바르지 않습니다.")
    ledger = state.get("ledger")
    if not isinstance(state.get("ledger_snapshots", {}), dict):
        raise ValidationError("통제원장 이력 구조가 올바르지 않습니다.")
    if any(not isinstance(v,dict) or k != v.get("as_of") for k,v in state.get("ledger_snapshots", {}).items()):
        raise ValidationError("통제원장 이력의 기준일과 저장 키가 다릅니다.")
    for ledger in ([ledger] if ledger is not None else []) + list(state.get("ledger_snapshots", {}).values()):
        if not isinstance(ledger, dict) or not isinstance(ledger.get("lines"), dict) or not isinstance(ledger.get("total"), dict):
            raise ValidationError("통제원장 구조가 올바르지 않습니다.")
        if not parse_date(ledger.get("as_of")) or parse_date(ledger["as_of"]).year != state["year"]:
            raise ValidationError("통제원장 기준일과 회계연도가 다릅니다.")
        if parse_date(ledger.get("period_start")) != date(state["year"], 1, 1):
            raise ValidationError("통제원장은 회계연도 1월 1일부터의 누계여야 합니다.")
        for field in ("current_budget", "allocated", "expense"):
            if type(ledger["total"].get(field)) is not int or any(type(x.get(field)) is not int for x in ledger["lines"].values()):
                raise ValidationError("통제원장 금액은 원 단위 정수여야 합니다.")
            if sum(x[field] for x in ledger["lines"].values()) != ledger["total"][field]:
                raise ValidationError("백업 통제원장 합계가 세목 합계와 다릅니다.")


def build_model(state):
    values, meta, semoks, timeline = {}, {}, {}, []
    for stage in STAGES:
        doc = state["stages"].get(stage)
        if not doc:
            continue
        before = values.copy()
        mode = doc.get("mode", "full")
        meta.update(doc["leaves"]); semoks.update(doc.get("semoks", {}))
        if mode == "full":
            values = {k: x["amount"] for k, x in doc["leaves"].items()}
        elif mode == "partial":
            values.update({k: x["amount"] for k, x in doc["leaves"].items()})
        elif mode == "delta":
            for k, x in doc["leaves"].items():
                values[k] = values.get(k, 0) + x["amount"]
        else:
            raise ValidationError("알 수 없는 추경 반영 방식입니다.")
        if any(v < 0 for v in values.values()) and not (doc.get("layout") == "erp_mixed_delta" and doc.get("review_acknowledged")):
            raise ValidationError(f"{stage} 반영 후 음수 예산이 발생합니다.")
        active = [meta[k] for k, v in values.items() if v]
        for a in active:
            if any(b["scope"] == a["scope"] and b["code"] == a["code"] and len(b["path"]) > len(a["path"]) and b["path"][:len(a["path"])] == a["path"] for b in active):
                raise ValidationError("기존 부기와 새 하위 부기가 동시에 예산을 가집니다. 전체 경정명세로 반영하거나 종전 부기를 명시적으로 감액해 주세요.")
        timeline.append({"stage": stage, "values": values.copy(), "total": sum(values.values()),
                         "delta": sum(values.values())-sum(before.values()) if stage != "본예산" else 0})
    leaves = {k: {**m, "amount": values.get(k, 0), "historical": k not in values} for k, m in meta.items()}
    return {"leaves": leaves, "semoks": semoks, "timeline": timeline, "total": sum(values.values()),
            "review_leaves": [x for x in leaves.values() if x["amount"] < 0]}


def prepare_amendment(state, label, doc, mapping=None):
    """Compare ERP changes with the preceding stage, keeping every signed row.

    Negative balances are source review items, never rounded, clamped or fixed
    by changing a sign. Acknowledgement permits a provisional source import.
    """
    result = copy.deepcopy(doc)
    if result.get("layout") != "erp_mixed_delta":
        return result
    result["errors"] = list(doc.get("parse_errors", doc.get("errors", [])))
    result["parse_errors"] = result["errors"].copy()
    if label not in STAGES[1:]:
        result["errors"].append("이 파일은 증감 부기가 있는 추경입니다. 제1회 또는 제2회 추경으로 등록하세요.")
        return result
    prior = copy.deepcopy(state)
    prior["stages"] = {s: state["stages"][s] for s in STAGES[:STAGES.index(label)] if s in state["stages"]}
    if len(prior["stages"]) != STAGES.index(label):
        result["errors"].append("본예산부터 앞선 차수를 순서대로 등록해 주세요.")
        return result
    model = build_model(prior)
    previous = model["leaves"]
    if result["previous_amount"] != model["total"]:
        result["errors"].append(f"원본 직전 예산 {result['previous_amount']:,}원과 등록된 직전 차수 {model['total']:,}원이 다릅니다.")
    for key, semok in result["semoks"].items():
        amount = sum(x["amount"] for x in previous.values() if x["scope"]+"|"+x["code"] == key)
        if amount != semok["previous_amount"]:
            result["errors"].append(f"{key}: 원본 직전 예산과 등록 부기 합계가 {amount-semok['previous_amount']:+,}원 다릅니다.")
    source = copy.deepcopy(doc.get("source_leaves", doc["leaves"]))
    result["source_leaves"] = source
    choices = dict(doc.get("mapping", {})) if mapping is None else dict(mapping)
    resolved, changes = {}, []
    for source_id, node in source.items():
        target = choices.get(source_id, source_id)
        reason = "동일 부기" if target in previous else "새 부기"
        # ERP renamed the one return entry from its facility-specific title to
        # a confirmed-return title. Restrict this suggestion to that exact code.
        if source_id not in choices and source_id not in previous and node["code"] == "802-11" and "반환금" in node["label"]:
            candidates = [k for k, x in previous.items() if x["scope"] == node["scope"] and x["code"] == node["code"] and "반환금" in x["label"] and not x["historical"]]
            if len(candidates) == 1:
                target = candidates[0]
                reason = "반환금 명칭 변경 · 같은 경비·세목의 유일한 부기에 연결"
        if target != source_id:
            old = previous.get(target)
            if not old or old["scope"] != node["scope"] or old["code"] != node["code"]:
                result["errors"].append(f"{node['label']}: 같은 경비·세목 안의 부기에만 연결할 수 있습니다.")
                continue
            reason = reason if reason.startswith("반환금") else "사용자가 연결한 부기"
        base = previous.get(target, node)
        components = copy.deepcopy(node.get("components", []))
        if target in resolved:
            resolved[target]["amount"] += node["amount"]
            resolved[target]["components"].extend(components)
        else:
            resolved[target] = {**node, "id": target, "path": base["path"], "label": base["label"], "components": components}
        choices[source_id] = target
        changes.append({"source_id": source_id, "target_id": target, "scope": node["scope"], "code": node["code"],
                        "source_label": node["label"], "label": base["label"], "delta": node["amount"], "match": reason,
                        "rows": [c["row"] for c in components] or [node["row"]]})
    after = {k: x["amount"] for k, x in previous.items()}
    for key, node in resolved.items():
        after[key] = after.get(key, 0) + node["amount"]
    for change in changes:
        change["before"] = previous.get(change["target_id"], {}).get("amount", 0)
        change["after"] = after[change["target_id"]]
    if sum(after.values()) != result["source_total"]:
        result["errors"].append("반영 후 전체 편성액이 원본 추경 총액과 다릅니다.")
    for key, semok in result["semoks"].items():
        total = sum(v for k, v in after.items() if (resolved.get(k) or previous[k])["scope"]+"|"+(resolved.get(k) or previous[k])["code"] == key)
        if total != semok["amount"]:
            result["errors"].append(f"{key}: 반영 후 세목 합계가 원본 경정액과 다릅니다.")
    result.update(leaves=resolved, mapping=choices, changes=changes,
                  review_items=[{"id": k, "label": (resolved.get(k) or previous[k])["label"], "amount": v,
                                 "scope": (resolved.get(k) or previous[k])["scope"], "code": (resolved.get(k) or previous[k])["code"]}
                                for k, v in after.items() if v < 0])
    return result


def apply_source(state, label, doc, filename, *, as_of=None, period_start=None, history_only=False):
    if label in STAGES and doc.get("layout") == "erp_mixed_delta":
        doc = prepare_amendment(state, label, doc)
        if not doc.get("errors") and doc.get("review_items") and not doc.get("review_acknowledged"):
            raise ValidationError("감액 후 음수인 부기가 있습니다. 원본 확인 필요 내역을 확인하고 잠정 반영에 동의해 주세요.")
    if doc.get("errors"):
        raise ValidationError("검증 오류를 해결한 뒤 반영해 주세요: " + "; ".join(doc["errors"][:3]))
    new = copy.deepcopy(state)
    doc = copy.deepcopy(doc)
    doc.update({"filename": filename, "imported_at": datetime.now().isoformat(timespec="seconds"), "year": state["year"]})
    if label in STAGES:
        idx = STAGES.index(label)
        if any(s not in new["stages"] for s in STAGES[:idx]):
            raise ValidationError("앞선 예산 차수부터 등록해야 차수별 증감을 정확히 계산할 수 있습니다.")
        if label == "본예산" and doc["mode"] != "full":
            raise ValidationError("본예산은 전체명세로 등록해 주세요.")
        # Replacing an earlier stage invalidates the assumptions of later stages.
        downstream = [s for s in STAGES[idx+1:] if s in new["stages"]]
        meaning = ("fingerprint", "mode", "sheet", "column", "scale", "leaves", "semoks")
        previous = new["stages"].get(label, {})
        if downstream and any(previous.get(k) != doc.get(k) for k in meaning):
            raise ValidationError("후속 추경이 등록되어 있습니다. 변경 이력에서 후속 추경 이전으로 복원한 뒤 순서대로 반영해 주세요.")
        new["stages"][label] = doc
        build_model(new)
    elif label == "지출통제원장":
        end, start = parse_date(as_of), parse_date(period_start)
        if not end or end.year != state["year"] or start != date(state["year"], 1, 1):
            raise ValidationError("해당 회계연도 1월 1일부터 조회종료일까지의 누계자료가 필요합니다.")
        old = new.get("ledger")
        if history_only and old and end.isoformat() >= old["as_of"]:
            raise ValidationError("과거 대조자료는 현재 조회종료일 이전 자료여야 합니다. 최신 자료는 일반 반영을 사용하세요.")
        if old and end.isoformat() < old["as_of"] and not history_only:
            raise ValidationError("현재 저장본보다 과거 기준일입니다. 변경 이력에서 해당 시점 자료를 확인해 주세요.")
        doc.update({"as_of": end.isoformat(), "period_start": start.isoformat()})
        snapshots = new.setdefault("ledger_snapshots", {})
        if old:
            snapshots[old["as_of"]] = copy.deepcopy(old)
        snapshots[end.isoformat()] = copy.deepcopy(doc)
        if not history_only:
            new["ledger"] = doc  # replace snapshot, never sum repeated imports
    else:
        raise ValidationError("알 수 없는 원천자료 구분입니다.")
    return new


def transaction_id(row):
    linked = row.get("원천연계키")
    if isinstance(linked, str) and linked.startswith(("id:", "erp:", "row:")):
        return linked
    explicit = row.get("거래ID") or row.get("transaction_id")
    if explicit:
        return "id:" + str(explicit)
    docid = row.get("지출결의번호") or row.get("지급명령번호") or row.get("문서번호")
    source_line = row.get("명세번호") or row.get("상세순번")
    if docid and source_line:
        return "erp:" + str(docid) + ":" + str(source_line)
    d = parse_date(row.get("집행일자"))
    scope_text = norm(row.get("사업구분", ""))
    scope_key = "일상경비" if "일상" in scope_text else "운영경비" if "수탁" in scope_text or "운영" in scope_text else ""
    official_code = expense_code(row)
    if not official_code and row.get("예산세목확인필요"):
        # Match the older special-import identity without treating its inferred
        # management label as an authoritative bookkeeping code.
        official_code = CATEGORY_CODES.get(row.get("관리항목"))
    official_code = official_code or norm(row.get("세목"))+"|"+norm(row.get("예산과목"))
    fields = [d.isoformat() if d else str(row.get("집행일자", "")), norm(row.get("적요")),
              won(row.get("집행금액", 0)), official_code, scope_key,
              norm(row.get("거래처", row.get("채주", ""))), str(docid or ""), str(row.get("중복구분", ""))]
    return "row:" + digest(json.dumps(fields, ensure_ascii=False).encode())


def merge_unique(old_rows, new_rows):
    result = list(old_rows)
    seen = {transaction_id(r) for r in old_rows}
    added = 0
    for row in new_rows:
        key = transaction_id(row)
        if key not in seen:
            result.append(copy.deepcopy(row)); seen.add(key); added += 1
    return sorted(result, key=lambda r: str(r.get("집행일자", "")), reverse=True), added


def expense_code(row):
    if row.get('업로드구분')=='수동집계조정' and row.get('예산세목확인필요'):
        return ''
    official, _origin, _raw = source_account(row)
    if official:
        return official
    if _origin in ("원본I열", "예산과목"):
        return ""
    if row.get("예산세목확인필요") and not row.get("세목") and not row.get("예산과목"):
        return ""
    for field in ("세목", "예산과목", "관리항목"):
        if str(row.get(field, "")) in CATEGORY_CODES:
            return CATEGORY_CODES[str(row[field])]
    return ""


def transaction_rows(state, daily):
    model = build_model(state)
    result, seen = [], set()
    cutoff = parse_date((state.get("ledger") or {}).get("as_of"))
    for original in daily:
        try:
            tid, amount = transaction_id(original), won(original.get("집행금액", 0))
        except ValidationError:
            result.append({"id": "bad:"+str(len(result)), "row": original, "amount": None, "status": "금액 확인", "allocations": [], "eligible": False})
            continue
        if tid in seen:
            continue
        seen.add(tid)
        dt = parse_date(original.get("집행일자"))
        code = expense_code(original)
        matching_scopes = {x["scope"] for x in model["leaves"].values() if x["code"] == code}
        scope_text = norm(original.get("사업구분", ""))
        scope = "일상경비" if "일상" in scope_text else "운영경비" if "수탁" in scope_text or "운영" in scope_text else ""
        if not scope and len(matching_scopes) == 1:
            scope = next(iter(matching_scopes))
        # Both branches can use the same code: require an explicit choice.
        override = state.get("overrides", {}).get(tid, {})
        scope, code = override.get("scope", scope), code or override.get("code", code)
        allocations = state["assignments"].get(tid, {}).get("allocations", [])
        reason = ""
        if not dt:
            reason = "지급일 확인"
        elif dt.year != state["year"]:
            reason = "다른 회계연도"
        elif cutoff and dt > cutoff:
            reason = "통제원장 기준일 이후"
        elif cutoff and original.get("일자근거") == "지급월(일자미상)" and dt.year == cutoff.year and dt.month == cutoff.month and cutoff.day < monthrange(cutoff.year, cutoff.month)[1]:
            reason = "지급일 확인"
        elif original.get("거래상태") in ("취소", "취소완료", "삭제", "미지급", "미집행", "반려"):
            reason = "취소·미지급 확인"
        elif not code or not scope:
            reason = "사업구분·세목 확인"
        elif not any(x["scope"] == scope and x["code"] == code for x in model["leaves"].values()):
            reason = "편성 부기 없음"
        valid = bool(allocations) and sum(won(a["amount"]) for a in allocations) == amount
        valid = valid and all(a["leaf"] in model["leaves"] and model["leaves"][a["leaf"]]["scope"] == scope and model["leaves"][a["leaf"]]["code"] == code for a in allocations)
        valid = valid and all(model["leaves"][a["leaf"]]["amount"] >= 0 for a in allocations)
        saved = state["assignments"].get(tid, {})
        if saved.get("source_signature") and saved["source_signature"] != transaction_signature(original):
            valid = False
        if saved.get("source_revision", 0) != original.get("원천수정번호", 0):
            valid = False
        result.append({"id": tid, "row": original, "amount": amount, "date": dt.isoformat() if dt else "",
                       "scope": scope, "code": code, "eligible": not reason,
                       "status": reason or ("분류완료" if valid else "재확인" if allocations else "미분류"),
                       "allocations": allocations if valid and not reason else []})
    return result


def classify(state, tx, scope, code, allocations, reason=""):
    if not tx.get("date") or parse_date(tx["date"]).year != state["year"]:
        raise ValidationError("지급일과 회계연도를 먼저 확인해 주세요.")
    original_code = expense_code(tx["row"])
    if original_code and code != original_code:
        raise ValidationError("원천에서 확인된 공식 세목은 부기 분류 중 변경할 수 없습니다. 원천자료에서 수정해 주세요.")
    model = build_model(state)
    cleaned, used = [], set()
    for allocation in allocations:
        key, amount = allocation.get("leaf"), won(allocation.get("amount"))
        leaf = model["leaves"].get(key)
        if not leaf or leaf["scope"] != scope or leaf["code"] != code:
            raise ValidationError("선택한 사업구분·세목의 실제 부기만 분류할 수 있습니다.")
        if leaf["amount"] < 0:
            raise ValidationError("감액 후 음수인 부기는 원본 확인이 필요하므로 분류 확정할 수 없습니다.")
        if key in used:
            raise ValidationError("같은 부기는 한 행으로 합쳐 주세요.")
        if amount and amount * tx["amount"] < 0:
            raise ValidationError("분할 금액은 원거래와 부호가 같아야 합니다.")
        cleaned.append({"leaf": key, "amount": amount}); used.add(key)
    if not cleaned or sum(x["amount"] for x in cleaned) != tx["amount"]:
        raise ValidationError("부기 분할 합계가 원거래 금액과 정확히 같아야 합니다.")
    if (scope, code) != (tx.get("scope"), tx.get("code")) and not reason.strip():
        raise ValidationError("사업구분 또는 세목을 확인·수정한 사유를 입력해 주세요.")
    new = copy.deepcopy(state)
    new["overrides"][tx["id"]] = {"scope": scope, "code": code, "reason": reason}
    new["assignments"][tx["id"]] = {"allocations": cleaned, "description": tx["row"].get("적요", ""),
                                     "updated_at": datetime.now().isoformat(timespec="seconds"), "reason": reason,
                                     "source_signature": transaction_signature(tx["row"]),
                                     "source_revision": tx["row"].get("원천수정번호", 0)}
    return new


def reconciliation(state, daily):
    model = build_model(state)
    txs = transaction_rows(state, daily)
    leaf_spent, linked, assigned = defaultdict(int), defaultdict(int), defaultdict(int)
    pending_count, pending_abs = defaultdict(int), defaultdict(int)
    for tx in txs:
        if tx["eligible"]:
            k = tx["scope"]+"|"+tx["code"]
            linked[k] += tx["amount"]
            if not tx["allocations"]:
                pending_count[k] += 1
                pending_abs[k] += abs(tx["amount"])
            for a in tx["allocations"]:
                leaf_spent[a["leaf"]] += a["amount"]
                assigned[k] += a["amount"]
    budgets = defaultdict(int)
    budget_review_keys = {x["scope"]+"|"+x["code"] for x in model.get("review_leaves", [])}
    for leaf in model["leaves"].values():
        budgets[leaf["scope"]+"|"+leaf["code"]] += leaf["amount"]
    ledger = (state.get("ledger") or {}).get("lines", {})
    keys = set(budgets) | set(ledger) | set(linked)
    rows = []
    for key in sorted(keys):
        scope, code = key.split("|")
        line = ledger.get(key)
        name = (line or model["semoks"].get(key, {})).get("name", code)
        official = line["expense"] if line else None
        current = line["current_budget"] if line else None
        rows.append({"key": key, "scope": scope, "code": code, "name": name, "budget": budgets[key],
                     "current": current, "official": official, "allocated": line.get("allocated") if line else None,
                     "gap": budgets[key]-current if current is not None else None,
                     "remaining": current-official if line else None, "linked": linked[key], "classified": assigned[key],
                     "unclassified": linked[key]-assigned[key], "unlinked": official-linked[key] if line else None,
                     "pending_count": pending_count[key], "pending_abs": pending_abs[key],
                     "budget_review": key in budget_review_keys,
                     "complete": bool(line) and pending_count[key] == 0 and official == linked[key] and key not in budget_review_keys,
                     "unexplained": official-assigned[key] if line else None})
    problems = [t for t in txs if t["status"] not in ("분류완료", "다른 회계연도", "통제원장 기준일 이후")]
    issues = {"pending_count": sum(pending_count.values()), "pending_abs": sum(pending_abs.values()),
              "unlinked_abs": sum(abs(r["unlinked"]) for r in rows if r["unlinked"] is not None),
              "unscoped_count": sum(not t["eligible"] for t in problems),
              "complete_count": sum(r["complete"] for r in rows), "semok_count": len(rows)}
    issues["all_complete"] = bool(ledger) and bool(rows) and issues["complete_count"] == len(rows) and not problems
    return {"model": model, "rows": rows, "transactions": txs, "leaf_spent": dict(leaf_spent), "issues": issues}


def recommendations(state, tx, candidates):
    # Suggestions only: previously confirmed exact descriptions, then scoped path tokens.
    matches = []
    text = norm(tx["row"].get("적요"))
    previous = {a["leaf"] for v in state["assignments"].values() if norm(v.get("description")) == text for a in v.get("allocations", [])}
    for leaf in candidates:
        tokens = [t for t in re.split(r"[^가-힣a-zA-Z0-9]+", leaf["label"]) if len(t) >= 2]
        score = 100 if leaf["id"] in previous else sum(len(t) for t in tokens if norm(t) in text)
        if score:
            matches.append((score, leaf["id"], "같은 적요의 이전 확정 분류" if score == 100 else "적요와 부기명 일치"))
    return sorted(matches, reverse=True)[:3]


def migrate_legacy_classifications(state, legacy, daily):
    """Recover only exact final-leaf matches after original budgets are imported."""
    mappings = legacy.get("classifications")
    if not isinstance(mappings, dict):
        raise ValidationError("V39/V43 예산 저장파일의 classifications 항목이 필요합니다.")
    model = build_model(state)
    if not model["leaves"]:
        raise ValidationError("기존 분류를 대조할 본예산 원본을 먼저 등록해 주세요.")
    new, imported, pending = copy.deepcopy(state), [], []
    for tx in transaction_rows(state, daily):
        row = tx["row"]
        if tx["amount"] is None:
            continue
        old39 = hashlib.md5("|".join(str(row.get(k, "")) for k in ("집행일자", "집행금액", "적요", "예산과목", "세목")).encode()).hexdigest()
        old43 = hashlib.sha1("|".join([str(row.get("집행일자", "")).strip(), str(row.get("적요", "")).strip(), str(tx["amount"]), str(row.get("세목", "")).strip(), str(row.get("예산과목", "")).strip()]).encode()).hexdigest()[:16]
        value = mappings.get(old39, mappings.get(old43))
        if not value or tx["id"] in new["assignments"]:
            continue
        if isinstance(value, dict):
            scope, code, label = value.get("scope", "일상경비"), value.get("semok_code", ""), value.get("bugi_name", "")
        elif isinstance(value, str) and len(value.split("|", 2)) == 3:
            scope, code, label = value.split("|", 2)
        else:
            pending.append({"transaction": tx["id"], "value": value}); continue
        leaves = [x for x in model["leaves"].values() if x["scope"] == scope and x["code"] == code and x["label"] == label]
        if len(leaves) == 1 and tx.get("date") and parse_date(tx["date"]).year == state["year"]:
            new = classify(new, tx, scope, code, [{"leaf": leaves[0]["id"], "amount": tx["amount"]}], "기존 확정 분류의 세목·전체 부기 경로 일치 확인")
            imported.append({"적요": row.get("적요", ""), "금액": tx["amount"], "부기": label})
        else:
            pending.append({"transaction": tx["id"], "scope": scope, "code": code, "label": label, "value": value})
    new["legacy_pending"] = pending
    return new, imported, pending


class BudgetStore:
    """Atomic state + original bytes + event history with optimistic concurrency."""
    def __init__(self, path="local_budget_v44.sqlite3"):
        self.path = str(path)
        with self.connect() as db:
            db.executescript("""CREATE TABLE IF NOT EXISTS states(year INTEGER PRIMARY KEY, revision INTEGER NOT NULL, payload TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, year INTEGER, revision INTEGER, at TEXT, action TEXT, payload TEXT);
            CREATE TABLE IF NOT EXISTS sources(hash TEXT PRIMARY KEY, filename TEXT, data BLOB);""")

    def connect(self):
        return self._connection()

    @contextmanager
    def _connection(self):
        db = sqlite3.connect(self.path, timeout=10)
        try:
            with db:
                yield db
        finally:
            db.close()

    def load(self, year):
        with self.connect() as db:
            row = db.execute("SELECT revision,payload FROM states WHERE year=?", (year,)).fetchone()
        return (json.loads(row[1]), row[0]) if row else (empty_state(year), 0)

    def save(self, state, expected_revision, action, source=None):
        validate_state(state)
        build_model(state)
        payload = json.dumps(state, ensure_ascii=False, allow_nan=False)
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT revision FROM states WHERE year=?", (state["year"],)).fetchone()
            actual = row[0] if row else 0
            if actual != expected_revision:
                raise ConflictError("다른 창에서 자료가 변경되었습니다. 화면을 새로고침한 뒤 다시 저장해 주세요.")
            revision = actual + 1
            if source:
                filename, raw = source
                db.execute("INSERT OR IGNORE INTO sources VALUES(?,?,?)", (digest(raw), filename, raw))
            db.execute("INSERT OR REPLACE INTO states VALUES(?,?,?)", (state["year"], revision, payload))
            db.execute("INSERT INTO events(year,revision,at,action,payload) VALUES(?,?,?,?,?)", (state["year"], revision,
                       datetime.now().isoformat(timespec="seconds"), action, payload))
        return revision

    def history(self, year):
        with self.connect() as db:
            return db.execute("SELECT revision,at,action FROM events WHERE year=? ORDER BY revision DESC", (year,)).fetchall()

    def version(self, year, revision):
        with self.connect() as db:
            row = db.execute("SELECT payload FROM events WHERE year=? AND revision=?", (year, revision)).fetchone()
        if not row:
            raise ValidationError("해당 이력이 없습니다.")
        return json.loads(row[0])

    def original(self, fingerprint):
        with self.connect() as db:
            return db.execute("SELECT filename,data FROM sources WHERE hash=?", (fingerprint,)).fetchone()

    def export(self, year):
        state, revision = self.load(year)
        hashes = {d["fingerprint"] for d in list(state["stages"].values()) + ([state["ledger"]] if state["ledger"] else [])}
        hashes.update(d["fingerprint"] for d in state.get("ledger_snapshots", {}).values())
        originals = []
        for h in hashes:
            source = self.original(h)
            if source:
                originals.append({"filename": source[0], "hash": h, "data": base64.b64encode(source[1]).decode()})
        return json.dumps({"format": "budget-v44", "state": state, "revision": revision, "originals": originals}, ensure_ascii=False).encode()

    def restore(self, raw, year, expected_revision):
        try:
            payload = json.loads(raw)
            state = payload["state"]
            validate_state(state)
            if payload.get("format") != "budget-v44" or state["year"] != year:
                raise ValidationError("선택 회계연도와 일치하는 V44 백업을 선택해 주세요.")
            for item in payload.get("originals", []):
                if digest(base64.b64decode(item["data"], validate=True)) != item["hash"]:
                    raise ValidationError("백업 원본의 검증값이 다릅니다.")
        except (KeyError, TypeError, ValueError, AttributeError) as exc:
            raise ValidationError(f"백업 확인 실패: {exc}") from exc
        # Saving imported state is one CAS transaction; original files follow in the same transaction below.
        build_model(state)
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT revision FROM states WHERE year=?", (year,)).fetchone()
            if (row[0] if row else 0) != expected_revision:
                raise ConflictError("다른 창에서 변경되었습니다. 새로고침해 주세요.")
            for item in payload.get("originals", []):
                db.execute("INSERT OR IGNORE INTO sources VALUES(?,?,?)", (item["hash"], item["filename"], base64.b64decode(item["data"])))
            serialized = json.dumps(state, ensure_ascii=False, allow_nan=False)
            revision = expected_revision+1
            db.execute("INSERT OR REPLACE INTO states VALUES(?,?,?)", (year, revision, serialized))
            db.execute("INSERT INTO events(year,revision,at,action,payload) VALUES(?,?,?,?,?)", (year, revision, datetime.now().isoformat(timespec="seconds"), "백업 복원", serialized))
        return revision
